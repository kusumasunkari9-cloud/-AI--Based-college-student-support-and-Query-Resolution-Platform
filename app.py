import difflib
import json
import math
import os
import re
import secrets
import sqlite3
import zipfile
from collections import Counter
from datetime import datetime, timezone
from functools import wraps
from pathlib import Path

import click
import pymupdf
import pytesseract
from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from dotenv import load_dotenv
from flask import (
    Flask,
    abort,
    flash,
    g,
    jsonify,
    redirect,
    render_template,
    request,
    session,
    url_for,
)
from PIL import Image
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename


BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.getenv("SECRET_KEY") or secrets.token_hex(32),
    DATABASE=str(BASE_DIR / "data" / "student_support.sqlite3"),
    UPLOAD_FOLDER=str(BASE_DIR / "uploads"),
    MAX_CONTENT_LENGTH=20 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.getenv("COOKIE_SECURE", "0") == "1",
)

ALLOWED_EXTENSIONS = {"pdf", "docx", "txt", "png", "jpg", "jpeg"}
STOP_WORDS = {
    "about", "after", "again", "also", "and", "are", "can", "could", "does",
    "for", "from", "have", "how", "into", "its", "need", "please", "should",
    "that", "the", "their", "them", "then", "there", "this", "those", "what",
    "when", "where", "which", "who", "will", "with", "would", "you", "your",
}
INTENT_KEYWORDS = {
    "Exams": ("exam", "examination", "hall ticket", "result", "revaluation", "backlog"),
    "Fees": ("fee", "fees", "tuition", "payment", "refund", "fine", "dues"),
    "Attendance": ("attendance", "absent", "present", "shortage", "condonation"),
    "Academics": ("course", "syllabus", "semester", "credit", "grading", "academic", "subject"),
    "Scholarships": ("scholarship", "stipend", "financial aid", "fellowship"),
    "Placements": ("placement", "internship", "career", "recruitment", "company"),
    "Library": ("library", "book", "borrow", "renewal", "reading room"),
    "Hostel": ("hostel", "dorm", "mess", "warden", "accommodation"),
    "Transport": ("transport", "bus", "route", "commute", "shuttle"),
}
PRIORITY_HIGH = ("urgent", "emergency", "deadline today", "locked out", "harassment", "safety")


def utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def valid_email(email):
    return re.fullmatch(r"[^@\s]+@[^@\s.]+(?:\.[^@\s.]+)+", email) is not None


def get_db():
    if "db" not in g:
        Path(app.config["DATABASE"]).parent.mkdir(parents=True, exist_ok=True)
        g.db = sqlite3.connect(app.config["DATABASE"])
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(_error=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


@app.errorhandler(413)
def request_too_large(_error):
    message = "The upload is larger than the 20 MB limit."
    if request.is_json:
        return jsonify({"error": message}), 413
    return render_template("error.html", status=413, message=message), 413


def init_db():
    db = get_db()
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY,
            email TEXT NOT NULL UNIQUE COLLATE NOCASE,
            name TEXT NOT NULL,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL CHECK (role IN ('student', 'admin')),
            department TEXT NOT NULL DEFAULT '',
            year TEXT NOT NULL DEFAULT '',
            semester TEXT NOT NULL DEFAULT '',
            courses TEXT NOT NULL DEFAULT '',
            preferences TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY,
            title TEXT NOT NULL,
            original_filename TEXT NOT NULL,
            stored_filename TEXT NOT NULL,
            owner_id INTEGER REFERENCES users(id) ON DELETE CASCADE,
            is_official INTEGER NOT NULL DEFAULT 0,
            academic_year TEXT NOT NULL DEFAULT '',
            semester TEXT NOT NULL DEFAULT '',
            department TEXT NOT NULL DEFAULT '',
            effective_date TEXT NOT NULL DEFAULT '',
            version TEXT NOT NULL DEFAULT '',
            content TEXT NOT NULL,
            uploaded_at TEXT NOT NULL,
            active INTEGER NOT NULL DEFAULT 1
        );
        CREATE INDEX IF NOT EXISTS idx_documents_title ON documents(title, department, uploaded_at);
        CREATE TABLE IF NOT EXISTS chunks (
            id INTEGER PRIMARY KEY,
            document_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
            page_number INTEGER,
            section TEXT NOT NULL DEFAULT '',
            content TEXT NOT NULL,
            embedding TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_chunks_document ON chunks(document_id);
        CREATE TABLE IF NOT EXISTS query_logs (
            id INTEGER PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            query TEXT NOT NULL,
            domain TEXT NOT NULL,
            confidence REAL NOT NULL,
            answer TEXT NOT NULL,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS tickets (
            id INTEGER PRIMARY KEY,
            user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            query_log_id INTEGER REFERENCES query_logs(id) ON DELETE SET NULL,
            subject TEXT NOT NULL,
            message TEXT NOT NULL,
            department TEXT NOT NULL,
            priority TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'Open',
            staff_response TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS feedback (
            id INTEGER PRIMARY KEY,
            query_log_id INTEGER NOT NULL REFERENCES query_logs(id) ON DELETE CASCADE,
            user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            rating INTEGER NOT NULL CHECK (rating IN (-1, 1)),
            created_at TEXT NOT NULL,
            UNIQUE(query_log_id, user_id)
        );
        CREATE TABLE IF NOT EXISTS policy_changes (
            id INTEGER PRIMARY KEY,
            old_document_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
            new_document_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
            summary TEXT NOT NULL,
            categories TEXT NOT NULL,
            affected_groups TEXT NOT NULL,
            effective_date TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL
        );
        """
    )
    document_columns = {row["name"] for row in db.execute("PRAGMA table_info(documents)").fetchall()}
    if "active" not in document_columns:
        db.execute("ALTER TABLE documents ADD COLUMN active INTEGER NOT NULL DEFAULT 1")
    db.commit()


@app.before_request
def load_logged_in_user():
    if request.method == "POST":
        submitted = request.form.get("csrf_token") or request.headers.get("X-CSRF-Token", "")
        expected = session.get("_csrf_token", "")
        if not expected or not secrets.compare_digest(submitted, expected):
            abort(400, "Invalid or missing CSRF token.")
    user_id = session.get("user_id")
    g.user = get_db().execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone() if user_id else None


@app.context_processor
def inject_csrf_token():
    if "_csrf_token" not in session:
        session["_csrf_token"] = secrets.token_urlsafe(32)
    return {"csrf_token": session["_csrf_token"]}


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if g.user is None:
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapped


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if g.user is None:
            return redirect(url_for("login"))
        if g.user["role"] != "admin":
            abort(403)
        return view(*args, **kwargs)
    return wrapped


def tokenize(text):
    return [word for word in re.findall(r"[a-z0-9]+", text.lower()) if len(word) > 1 and word not in STOP_WORDS]


def vectorize(text):
    return dict(Counter(tokenize(text)))


def cosine_similarity(query_vector, document_vector):
    if not query_vector or not document_vector:
        return 0.0
    common = set(query_vector).intersection(document_vector)
    dot = sum(query_vector[word] * document_vector[word] for word in common)
    query_norm = math.sqrt(sum(value * value for value in query_vector.values()))
    document_norm = math.sqrt(sum(value * value for value in document_vector.values()))
    return dot / (query_norm * document_norm) if query_norm and document_norm else 0.0


def detect_intent(query):
    normalized = query.lower()
    scores = {
        domain: sum(1 for phrase in phrases if phrase in normalized)
        for domain, phrases in INTENT_KEYWORDS.items()
    }
    domain = max(scores, key=scores.get) if any(scores.values()) else "General"
    entities = {}
    semester = re.search(r"\b(?:semester|sem)\s*[-:]?\s*(\d{1,2}|[ivx]+)\b", normalized)
    if semester:
        entities["semester"] = semester.group(1).upper()
    exam = re.search(r"\b(mid[- ]?term|end[- ]?term|supplementary|regular|final)\b", normalized)
    if exam:
        entities["exam_type"] = exam.group(1)
    date = re.search(r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2})\b", query)
    if date:
        entities["date"] = date.group(0)
    course = re.search(r"\b[A-Z]{2,5}[- ]?\d{2,4}\b", query)
    if course:
        entities["course"] = course.group(0)
    for value, pattern in (
        ("department", r"\b(?:department|dept)\s+(?:of\s+)?([A-Za-z][A-Za-z &-]{1,35})"),
        ("document", r"\b([\w -]{2,50}\.(?:pdf|docx|txt))\b"),
    ):
        match = re.search(pattern, query, re.IGNORECASE)
        if match:
            entities[value] = match.group(1).strip()
    return domain, entities


def split_chunks(text, page_number=None, section=""):
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + 1000, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + 650, end)
            if boundary > start:
                end = boundary
        part = text[start:end].strip()
        if part:
            chunks.append((page_number, section, part))
        if end >= len(text):
            break
        start = max(end - 120, start + 1)
    return chunks


def extract_document(path, extension):
    extracted = []
    if extension == "pdf":
        try:
            pdf = pymupdf.open(path)
        except pymupdf.FileDataError as exc:
            raise ValueError("The selected file is not a readable PDF.") from exc
        try:
            for page_index, page in enumerate(pdf, 1):
                page_text = page.get_text("text").strip()
                if not page_text:
                    try:
                        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                        image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                        page_text = pytesseract.image_to_string(image).strip()
                    except pytesseract.TesseractNotFoundError as exc:
                        raise ValueError(
                            "This scanned PDF needs Tesseract OCR. Install Tesseract and add it to PATH, then retry."
                        ) from exc
                extracted.append((page_index, page_text))
        finally:
            pdf.close()
    elif extension == "docx":
        try:
            document = DocxDocument(path)
        except (zipfile.BadZipFile, KeyError, PackageNotFoundError) as exc:
            raise ValueError("The selected file is not a readable DOCX document.") from exc
        lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        extracted.append((None, "\n".join(lines)))
    elif extension == "txt":
        try:
            extracted.append((None, Path(path).read_text(encoding="utf-8-sig")))
        except UnicodeDecodeError as exc:
            raise ValueError("The TXT document must use UTF-8 text encoding.") from exc
    elif extension in {"png", "jpg", "jpeg"}:
        try:
            with Image.open(path) as image:
                extracted.append((1, pytesseract.image_to_string(image).strip()))
        except pytesseract.TesseractNotFoundError as exc:
            raise ValueError("Image OCR needs Tesseract installed and available on PATH.") from exc
    content = "\n\n".join(text for _, text in extracted if text).strip()
    if not content:
        raise ValueError("No readable text was extracted from this document.")
    chunks = []
    for page_number, page_text in extracted:
        current_section = f"Page {page_number}" if page_number else "Document"
        lines = [line.strip() for line in page_text.splitlines() if line.strip()]
        block = []
        for line in lines:
            if len(line) < 100 and (line.isupper() or line.endswith(":")):
                if block:
                    chunks.extend(split_chunks("\n".join(block), page_number, current_section))
                    block = []
                current_section = line.rstrip(":")
            else:
                block.append(line)
        if block:
            chunks.extend(split_chunks("\n".join(block), page_number, current_section))
    return content, chunks


def infer_document_metadata(content, filename):
    source = f"{filename}\n{content}"
    inferred = {}
    academic_year = re.search(r"\b(20\d{2}\s*[-/]\s*(?:20)?\d{2})\b", source)
    if academic_year:
        inferred["academic_year"] = re.sub(r"\s+", "", academic_year.group(1)).replace("/", "-")
    semester = re.search(r"\b(?:semester|sem(?:ester)?)\s*[-:]?\s*(\d{1,2}|[ivx]+)\b", source, re.IGNORECASE)
    if semester:
        inferred["semester"] = semester.group(1).upper()
    effective_date = re.search(r"\beffective\s+(?:from|date)\s*:?\s*(\d{4}-\d{2}-\d{2})\b", source, re.IGNORECASE)
    if effective_date:
        inferred["effective_date"] = effective_date.group(1)
    version = re.search(r"\b(?:version|ver\.?|v)\s*[:.]?\s*(\d+(?:\.\d+)*|\d{4})\b", filename, re.IGNORECASE)
    if version:
        inferred["version"] = version.group(1)
    department = re.search(
        r"(?im)^\s*(?:department|dept\.?|school|faculty)\s+(?:of\s+)?([A-Z][A-Za-z &-]{2,45})\s*$",
        content,
    )
    if department:
        inferred["department"] = department.group(1).strip()
    return inferred


def add_document(file_storage, owner_id, is_official, metadata):
    filename = secure_filename(file_storage.filename or "")
    extension = Path(filename).suffix.lower().lstrip(".")
    if not filename or extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Choose a PDF, DOCX, TXT, PNG, JPG, or JPEG file.")
    stored_name = f"{secrets.token_hex(16)}.{extension}"
    upload_folder = Path(app.config["UPLOAD_FOLDER"])
    upload_folder.mkdir(parents=True, exist_ok=True)
    path = upload_folder / stored_name
    file_storage.save(path)
    try:
        content, chunks = extract_document(path, extension)
        inferred_metadata = infer_document_metadata(content, filename)
        db = get_db()
        title = (metadata.get("title") or Path(filename).stem).strip()[:180]
        now = utc_now()
        cursor = db.execute(
            """INSERT INTO documents
               (title, original_filename, stored_filename, owner_id, is_official, academic_year,
                semester, department, effective_date, version, content, uploaded_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                title, filename, stored_name, owner_id, int(is_official),
                (metadata.get("academic_year") or inferred_metadata.get("academic_year", "")).strip()[:40],
                (metadata.get("semester") or inferred_metadata.get("semester", "")).strip()[:40],
                (metadata.get("department") or inferred_metadata.get("department", "")).strip()[:100],
                (metadata.get("effective_date") or inferred_metadata.get("effective_date", "")).strip()[:40],
                (metadata.get("version") or inferred_metadata.get("version", "")).strip()[:60],
                content, now,
            ),
        )
        document_id = cursor.lastrowid
        db.executemany(
            "INSERT INTO chunks (document_id, page_number, section, content, embedding) VALUES (?, ?, ?, ?, ?)",
            [(document_id, page, section, chunk, json.dumps(vectorize(chunk))) for page, section, chunk in chunks],
        )
        db.commit()
        if is_official:
            detect_policy_change(document_id)
        return document_id
    except Exception:
        path.unlink(missing_ok=True)
        raise


def detect_policy_change(new_document_id):
    db = get_db()
    new_doc = db.execute("SELECT * FROM documents WHERE id = ?", (new_document_id,)).fetchone()
    old_doc = db.execute(
        """SELECT * FROM documents WHERE id != ? AND is_official = 1
           AND lower(title) = lower(?) AND lower(department) = lower(?) AND lower(semester) = lower(?)
           ORDER BY uploaded_at DESC LIMIT 1""",
        (new_document_id, new_doc["title"], new_doc["department"], new_doc["semester"]),
    ).fetchone()
    if old_doc is None:
        return
    old_text, new_text = old_doc["content"].splitlines(), new_doc["content"].splitlines()
    diff = list(difflib.unified_diff(old_text, new_text, fromfile="previous", tofile="new", lineterm=""))
    if not diff:
        return
    summary = "\n".join(diff[:80])
    combined = " ".join(diff).lower()
    categories = [
        category for category, terms in {
            "Attendance": ("attendance", "absent", "condonation"),
            "Examinations": ("exam", "examination", "hall ticket"),
            "Fees": ("fee", "fees", "refund", "fine"),
            "Deadlines": ("deadline", "last date", "due date", "effective"),
            "Eligibility": ("eligibility", "eligible", "minimum", "criteria"),
            "Academic regulations": ("credit", "grading", "semester", "academic"),
        }.items() if any(term in combined for term in terms)
    ]
    groups = []
    if new_doc["department"]:
        groups.append(new_doc["department"])
    if new_doc["semester"]:
        groups.append(f"Semester {new_doc['semester']}")
    if not groups:
        groups.append("All students")
    db.execute(
        """INSERT INTO policy_changes
           (old_document_id, new_document_id, summary, categories, affected_groups, effective_date, created_at)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        (old_doc["id"], new_document_id, summary, ", ".join(categories) or "Policy update",
         ", ".join(groups), new_doc["effective_date"], utc_now()),
    )
    db.commit()


def retrieve_evidence(user, query, limit=5):
    db = get_db()
    candidates = db.execute(
        """SELECT c.*, d.title, d.is_official, d.owner_id, d.department, d.semester,
                  d.academic_year, d.version, d.effective_date
           FROM chunks c JOIN documents d ON d.id = c.document_id
           WHERE (d.is_official = 1 AND d.active = 1) OR d.owner_id = ?
           ORDER BY d.is_official DESC, d.uploaded_at DESC""",
        (user["id"],),
    ).fetchall()
    query_entities = detect_intent(query)[1]
    profile_context_used = []
    retrieval_query = query
    if user["courses"] and any(term in query.lower() for term in ("my course", "my subject", "my classes")):
        retrieval_query += " " + user["courses"]
        profile_context_used.append(f"Courses: {user['courses']}")
    query_vector = vectorize(retrieval_query)
    vectors = [json.loads(row["embedding"]) for row in candidates]
    document_frequency = Counter(
        token for vector in vectors for token in vector.keys()
    )
    document_count = max(len(vectors), 1)
    inverse_document_frequency = {
        token: math.log((document_count + 1) / (frequency + 1)) + 1
        for token, frequency in document_frequency.items()
    }
    weighted_query = {
        token: frequency * inverse_document_frequency.get(token, 1.0)
        for token, frequency in query_vector.items()
    }
    scored = []
    for row, vector in zip(candidates, vectors):
        weighted_document = {
            token: frequency * inverse_document_frequency.get(token, 1.0)
            for token, frequency in vector.items()
        }
        score = cosine_similarity(weighted_query, weighted_document)
        if row["is_official"]:
            score *= 1.08
        query_department = (query_entities.get("department") or "").strip().lower()
        student_department = (user["department"] or "").strip().lower()
        target_department = query_department or student_department
        if target_department and row["department"]:
            score *= 1.15 if target_department == row["department"].strip().lower() else 0.4
            if student_department and not query_department:
                label = f"Department: {user['department']} (source ranking)"
                if label not in profile_context_used:
                    profile_context_used.append(label)
        query_semester = str(query_entities.get("semester") or "").lower().replace("semester", "").strip()
        profile_semester = (user["semester"] or "").lower().replace("semester", "").strip()
        student_semester = query_semester or profile_semester
        document_semester = (row["semester"] or "").lower().replace("semester", "").strip()
        if student_semester and document_semester and document_semester not in {"all", "all semesters"}:
            score *= 1.15 if student_semester == document_semester else 0.4
            if profile_semester and not query_semester:
                label = f"Semester: {user['semester']} (source ranking)"
                if label not in profile_context_used:
                    profile_context_used.append(label)
        scored.append((score, row))
    scored.sort(key=lambda item: item[0], reverse=True)
    selected = []
    per_document = Counter()
    for score, row in scored:
        if score <= 0 or per_document[row["document_id"]] >= 2:
            continue
        selected.append((score, row))
        per_document[row["document_id"]] += 1
        if len(selected) >= limit:
            break
    return selected, profile_context_used


def make_citation(row):
    location = f"page {row['page_number']}" if row["page_number"] else ""
    if row["section"] and row["section"] != f"Page {row['page_number']}":
        location = f"{location}, section {row['section']}" if location else f"section {row['section']}"
    version = f", version {row['version']}" if row["version"] else ""
    return f"{row['title']}{version}{f' ({location})' if location else ''}"


def create_ticket(user_id, query, domain, query_log_id=None):
    db = get_db()
    priority = "High" if any(word in query.lower() for word in PRIORITY_HIGH) else "Normal"
    department = "Student Services" if domain == "General" else domain
    now = utc_now()
    db.execute(
        """INSERT INTO tickets
           (user_id, query_log_id, subject, message, department, priority, status, created_at, updated_at)
           VALUES (?, ?, ?, ?, ?, ?, 'Open', ?, ?)""",
        (user_id, query_log_id, query[:180], query, department, priority, now, now),
    )
    db.commit()
    return db.execute("SELECT last_insert_rowid()").fetchone()[0]


def answer_query(user, query, explain=False):
    domain, entities = detect_intent(query)
    human_action_terms = ("please contact", "please call", "please review", "please correct",
                          "please update", "reissue", "resolve my", "complaint", "speak to staff")
    if any(term in query.lower() for term in human_action_terms):
        answer = "This request needs a staff action rather than a document-only answer. I have routed it to the appropriate support team."
        db = get_db()
        cursor = db.execute(
            "INSERT INTO query_logs (user_id, query, domain, confidence, answer, created_at) VALUES (?, ?, ?, ?, ?, ?)",
            (user["id"], query, domain, 0.0, answer, utc_now()),
        )
        db.commit()
        ticket_id = create_ticket(user["id"], query, domain, cursor.lastrowid)
        return {
            "answer": answer, "citations": [], "confidence": 0.0,
            "domain": domain, "entities": entities, "ticket_id": ticket_id,
            "query_log_id": cursor.lastrowid, "context_used": [],
        }
    evidence, profile_context_used = retrieve_evidence(user, query)
    best_score = evidence[0][0] if evidence else 0.0
    confidence = min(best_score, 1.0)
    threshold = 0.16
    if not evidence or confidence < threshold:
        answer = "I could not verify this from the available college documents or your uploaded documents. I have opened a support ticket so staff can help."
        confidence = round(confidence, 3)
        db = get_db()
        cursor = db.execute(
            "INSERT INTO query_logs (user_id, query, domain, confidence, answer, created_at) VALUES (?, ?, ?, ?, ?, ?)",
            (user["id"], query, domain, confidence, answer, utc_now()),
        )
        db.commit()
        ticket_id = create_ticket(user["id"], query, domain, cursor.lastrowid)
        return {
            "answer": answer, "citations": [], "confidence": confidence,
            "domain": domain, "entities": entities, "ticket_id": ticket_id,
            "query_log_id": cursor.lastrowid, "context_used": [],
        }
    sources = []
    snippets = []
    for score, row in evidence:
        citation = make_citation(row)
        if citation not in sources:
            sources.append(citation)
            snippets.append(f"• {row['content'].strip()} [{citation}]")
    context_used = profile_context_used
    if g.user["courses"]:
        courses = [course.strip() for course in g.user["courses"].split(",") if course.strip()]
        matched_courses = [course for course in courses if course.lower() in query.lower()]
        if matched_courses and not any(item.startswith("Courses:") for item in context_used):
            context_used.append(f"Course: {', '.join(matched_courses)}")
    answer_preferences = (g.user["preferences"] or "").lower()
    excerpt_limit = 2 if "concise" in answer_preferences else 4
    answer = "Here is what the available evidence says:\n\n" + "\n\n".join(snippets[:excerpt_limit])
    if "concise" in answer_preferences:
        context_used.append("Answer preference: concise")
    if context_used:
        answer += "\n\nYour saved profile context: " + "; ".join(context_used) + "."
    if explain:
        answer += "\n\nWhy this applies: the answer is based on the cited passages retrieved for your question. " \
                  "Profile context listed above was used only as additional student context; it does not replace a college rule."
    db = get_db()
    cursor = db.execute(
        "INSERT INTO query_logs (user_id, query, domain, confidence, answer, created_at) VALUES (?, ?, ?, ?, ?, ?)",
        (user["id"], query, domain, round(confidence, 3), answer, utc_now()),
    )
    db.commit()
    return {
        "answer": answer, "citations": sources, "confidence": round(confidence, 3),
        "domain": domain, "entities": entities, "ticket_id": None,
        "query_log_id": cursor.lastrowid, "context_used": context_used,
    }


@app.route("/")
def index():
    if g.user is None:
        return render_template("login.html")
    db = get_db()
    tickets = db.execute(
        "SELECT * FROM tickets WHERE user_id = ? ORDER BY updated_at DESC LIMIT 10",
        (g.user["id"],),
    ).fetchall()
    documents = db.execute(
        """SELECT id, title, original_filename, is_official, academic_year, semester,
                  department, effective_date, version, uploaded_at
           FROM documents WHERE (is_official = 1 AND active = 1) OR owner_id = ? ORDER BY uploaded_at DESC LIMIT 20""",
        (g.user["id"],),
    ).fetchall()
    if g.user["role"] == "admin":
        metrics = {
            "documents": db.execute("SELECT count(*) FROM documents WHERE is_official = 1").fetchone()[0],
            "students": db.execute("SELECT count(*) FROM users WHERE role = 'student'").fetchone()[0],
            "open_tickets": db.execute("SELECT count(*) FROM tickets WHERE status != 'Closed'").fetchone()[0],
            "unanswered": db.execute("SELECT count(*) FROM query_logs WHERE confidence < 0.16").fetchone()[0],
            "questions": db.execute("SELECT count(*) FROM query_logs").fetchone()[0],
            "avg_confidence": db.execute("SELECT coalesce(avg(confidence), 0) FROM query_logs").fetchone()[0],
        }
        feedback_stats = db.execute(
            """SELECT count(*) AS total, coalesce(sum(CASE WHEN rating = 1 THEN 1 ELSE 0 END), 0) AS helpful
               FROM feedback"""
        ).fetchone()
        metrics["feedback_total"] = feedback_stats["total"]
        metrics["feedback_helpful"] = feedback_stats["helpful"]
        frequent = db.execute(
            """SELECT query, domain, count(*) AS asked, avg(confidence) AS confidence
               FROM query_logs GROUP BY lower(query) ORDER BY asked DESC LIMIT 8"""
        ).fetchall()
        changes = db.execute(
            """SELECT pc.*, old.title AS title, old.version AS old_version,
                      new.version AS new_version, new.department AS department
               FROM policy_changes pc
               JOIN documents old ON old.id = pc.old_document_id
               JOIN documents new ON new.id = pc.new_document_id
               ORDER BY pc.created_at DESC LIMIT 10"""
        ).fetchall()
        all_tickets = db.execute(
            """SELECT t.*, u.name AS student_name, u.email AS student_email
               FROM tickets t JOIN users u ON u.id = t.user_id
               ORDER BY CASE t.priority WHEN 'High' THEN 0 ELSE 1 END, t.updated_at DESC LIMIT 30"""
        ).fetchall()
        low_queries = db.execute(
            """SELECT q.*, u.name AS student_name FROM query_logs q
               JOIN users u ON u.id = q.user_id WHERE q.confidence < 0.16
               ORDER BY q.created_at DESC LIMIT 12"""
        ).fetchall()
        managed_documents = db.execute(
            """SELECT id, title, version, department, effective_date, active, original_filename
               FROM documents WHERE is_official = 1 ORDER BY uploaded_at DESC LIMIT 50"""
        ).fetchall()
    else:
        metrics = frequent = changes = all_tickets = low_queries = []
        managed_documents = []
    comparison_groups = {}
    for document in db.execute(
        """SELECT id, title, department, semester, version, uploaded_at FROM documents
           WHERE is_official = 1 AND active = 1 ORDER BY uploaded_at, id"""
    ).fetchall():
        key = (document["title"].lower(), document["department"].lower(), document["semester"].lower())
        group = comparison_groups.setdefault(key, {
            "title": document["title"], "department": document["department"],
            "semester": document["semester"], "versions": [],
        })
        group["versions"].append({
            "id": document["id"], "version": document["version"],
            "uploaded_at": document["uploaded_at"],
        })
    docs_for_compare = [group for group in comparison_groups.values() if len(group["versions"]) > 1]
    return render_template(
        "dashboard.html", tickets=tickets, documents=documents, metrics=metrics,
        frequent=frequent, changes=changes, all_tickets=all_tickets, low_queries=low_queries,
        docs_for_compare=docs_for_compare, managed_documents=managed_documents,
    )


@app.route("/register", methods=["POST"])
def register():
    email = request.form.get("email", "").strip().lower()
    name = request.form.get("name", "").strip()
    password = request.form.get("password", "")
    if not valid_email(email) or not name or len(password) < 12:
        flash("Enter your name and email, and choose a password with at least 12 characters.", "error")
        return redirect(url_for("index"))
    try:
        cursor = get_db().execute(
            """INSERT INTO users (email, name, password_hash, role, created_at)
               VALUES (?, ?, ?, 'student', ?)""",
            (email, name[:100], generate_password_hash(password), utc_now()),
        )
        get_db().commit()
    except sqlite3.IntegrityError:
        flash("An account with that email already exists.", "error")
        return redirect(url_for("index"))
    session.clear()
    session["user_id"] = cursor.lastrowid
    return redirect(url_for("index"))


@app.route("/login", methods=["POST"])
def login():
    email = request.form.get("email", "").strip()
    password = request.form.get("password", "")
    user = get_db().execute("SELECT * FROM users WHERE email = ?", (email,)).fetchone()
    if user is None or not check_password_hash(user["password_hash"], password):
        flash("Email or password was not recognized.", "error")
        return redirect(url_for("index"))
    session.clear()
    session["user_id"] = user["id"]
    return redirect(url_for("index"))


@app.route("/logout", methods=["POST"])
@login_required
def logout():
    session.clear()
    return redirect(url_for("index"))


@app.route("/profile", methods=["POST"])
@login_required
def update_profile():
    get_db().execute(
        """UPDATE users SET name = ?, department = ?, year = ?, semester = ?, courses = ?, preferences = ?
           WHERE id = ?""",
        (
            request.form.get("name", "").strip()[:100],
            request.form.get("department", "").strip()[:100],
            request.form.get("year", "").strip()[:20],
            request.form.get("semester", "").strip()[:40],
            request.form.get("courses", "").strip()[:500],
            request.form.get("preferences", "").strip()[:500],
            g.user["id"],
        ),
    )
    get_db().commit()
    flash("Your profile was updated.", "success")
    return redirect(url_for("index"))


@app.route("/ask", methods=["POST"])
@login_required
def ask():
    payload = request.get_json(silent=True) or request.form
    query = str(payload.get("query", "")).strip()
    if not query or len(query) > 2000:
        return jsonify({"error": "Enter a question of 1–2,000 characters."}), 400
    result = answer_query(g.user, query, str(payload.get("explain", "")).lower() in {"1", "true", "yes"})
    return jsonify(result)


@app.route("/feedback", methods=["POST"])
@login_required
def submit_feedback():
    payload = request.get_json(silent=True) or request.form
    try:
        query_id, rating = int(payload.get("query_log_id")), int(payload.get("rating"))
    except (TypeError, ValueError):
        return jsonify({"error": "Invalid feedback."}), 400
    if rating not in (-1, 1):
        return jsonify({"error": "Invalid feedback rating."}), 400
    db = get_db()
    owned = db.execute("SELECT id FROM query_logs WHERE id = ? AND user_id = ?", (query_id, g.user["id"])).fetchone()
    if owned is None:
        abort(404)
    db.execute(
        """INSERT INTO feedback (query_log_id, user_id, rating, created_at) VALUES (?, ?, ?, ?)
           ON CONFLICT(query_log_id, user_id) DO UPDATE SET rating = excluded.rating, created_at = excluded.created_at""",
        (query_id, g.user["id"], rating, utc_now()),
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/documents/upload", methods=["POST"])
@login_required
def upload_personal_document():
    file_storage = request.files.get("document")
    if not file_storage:
        flash("Choose a document to upload.", "error")
    else:
        try:
            add_document(file_storage, g.user["id"], False, {"title": request.form.get("title", "")})
            flash("Your document was indexed and is only visible to your account.", "success")
        except (ValueError, OSError) as exc:
            flash(str(exc), "error")
    return redirect(url_for("index"))


@app.route("/tickets", methods=["POST"])
@login_required
def create_student_ticket():
    subject = request.form.get("subject", "").strip()
    message = request.form.get("message", "").strip()
    if not subject or not message or len(message) > 3000:
        flash("Enter a short subject and a message of up to 3,000 characters.", "error")
        return redirect(url_for("index"))
    domain, _ = detect_intent(f"{subject} {message}")
    ticket_id = create_ticket(g.user["id"], f"{subject}\n{message}", domain)
    db = get_db()
    db.execute(
        "UPDATE tickets SET subject = ?, message = ? WHERE id = ?",
        (subject[:180], message, ticket_id),
    )
    db.commit()
    flash(f"Support ticket #{ticket_id} was created.", "success")
    return redirect(url_for("index"))


@app.route("/admin/documents", methods=["POST"])
@admin_required
def upload_official_document():
    file_storage = request.files.get("document")
    if not file_storage:
        flash("Choose an official document to upload.", "error")
        return redirect(url_for("index"))
    metadata = {key: request.form.get(key, "") for key in (
        "title", "academic_year", "semester", "department", "effective_date", "version",
    )}
    try:
        add_document(file_storage, None, True, metadata)
        flash("Official document indexed. A change report was created if a previous version exists.", "success")
    except (ValueError, OSError) as exc:
        flash(str(exc), "error")
    return redirect(url_for("index"))


@app.route("/admin/documents/<int:document_id>/archive", methods=["POST"])
@admin_required
def toggle_document_archive(document_id):
    db = get_db()
    document = db.execute(
        "SELECT id, active FROM documents WHERE id = ? AND is_official = 1", (document_id,)
    ).fetchone()
    if document is None:
        abort(404)
    new_state = 0 if document["active"] else 1
    db.execute("UPDATE documents SET active = ? WHERE id = ?", (new_state, document_id))
    db.commit()
    flash(
        "Official document restored." if new_state
        else "Official document archived; it is no longer used for student answers.",
        "success",
    )
    return redirect(url_for("index"))


def compare_documents(old_id, new_id, student=None):
    db = get_db()
    old_doc = db.execute("SELECT * FROM documents WHERE id = ? AND is_official = 1 AND active = 1", (old_id,)).fetchone()
    new_doc = db.execute("SELECT * FROM documents WHERE id = ? AND is_official = 1 AND active = 1", (new_id,)).fetchone()
    if old_doc is None or new_doc is None:
        abort(404)
    if (
        old_doc["title"].lower() != new_doc["title"].lower()
        or old_doc["department"].lower() != new_doc["department"].lower()
        or old_doc["semester"].lower() != new_doc["semester"].lower()
        or old_id == new_id
    ):
        return {"error": "Select two different official versions with matching title, department, and semester."}, 400
    if old_id > new_id:
        return {"error": "Select the earlier version first and the later version second."}, 400
    old_lines = [line.strip() for line in old_doc["content"].splitlines() if line.strip()]
    new_lines = [line.strip() for line in new_doc["content"].splitlines() if line.strip()]
    old_chunks = db.execute(
        """SELECT c.*, d.title, d.version FROM chunks c JOIN documents d ON d.id = c.document_id
           WHERE c.document_id = ?""", (old_id,),
    ).fetchall()
    new_chunks = db.execute(
        """SELECT c.*, d.title, d.version FROM chunks c JOIN documents d ON d.id = c.document_id
           WHERE c.document_id = ?""", (new_id,),
    ).fetchall()

    def cite_matching_chunk(text, chunks, document):
        if not text or not chunks:
            return f"{document['title']} version {document['version'] or 'current'}"
        vector = vectorize(text)
        best = max(chunks, key=lambda chunk: cosine_similarity(vector, json.loads(chunk["embedding"])))
        citation = make_citation(best)
        return f"{citation}, version {document['version'] or 'current'}"

    matcher = difflib.SequenceMatcher(a=old_lines, b=new_lines)
    changes = []
    keywords = {
        "Attendance": ("attendance", "absent", "condonation"),
        "Examinations": ("exam", "examination", "hall ticket"),
        "Fees": ("fee", "fees", "refund", "fine"),
        "Deadlines": ("deadline", "due date", "last date"),
        "Eligibility": ("eligibility", "eligible", "criteria"),
        "Academic regulations": ("credit", "grading", "semester"),
    }
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag != "equal":
            old_text, new_text = " ".join(old_lines[i1:i2]), " ".join(new_lines[j1:j2])
            if old_text or new_text:
                changes.append({
                    "type": tag, "old": old_text, "new": new_text,
                    "citations": [
                        cite_matching_chunk(old_text, old_chunks, old_doc),
                        cite_matching_chunk(new_text, new_chunks, new_doc),
                    ],
                })
    joined = " ".join(item["old"] + " " + item["new"] for item in changes).lower()
    categories = [name for name, terms in keywords.items() if any(term in joined for term in terms)]
    affected = True
    reason = "The updated rule is not specific to a department or semester in the document metadata."
    if student is not None:
        if new_doc["department"] and student["department"]:
            affected = new_doc["department"].lower() == student["department"].lower()
            reason = "Department match" if affected else f"Document applies to {new_doc['department']}, not your saved department."
        if new_doc["semester"] and student["semester"]:
            semester_value = str(student["semester"]).lower().replace("semester", "").strip()
            semester_affected = semester_value == new_doc["semester"].lower().replace("semester", "").strip()
            affected = affected and semester_affected
            reason += "; semester match" if semester_affected else f"; this change is for semester {new_doc['semester']}."
    return {
        "title": old_doc["title"], "old_version": old_doc["version"] or "Previous version",
        "new_version": new_doc["version"] or "Current version",
        "effective_date": new_doc["effective_date"] or "Not specified",
        "categories": categories or ["Other policy changes"], "changes": changes[:40],
        "affected": affected, "reason": reason,
        "citations": [
            f"{old_doc['title']} version {old_doc['version'] or 'previous'}",
            f"{new_doc['title']} version {new_doc['version'] or 'current'}",
        ],
    }, 200


@app.route("/compare", methods=["POST"])
@login_required
def compare_route():
    payload = request.get_json(silent=True) or request.form
    try:
        old_id, new_id = int(payload.get("old_id")), int(payload.get("new_id"))
    except (TypeError, ValueError):
        return jsonify({"error": "Select both document versions."}), 400
    result, status = compare_documents(old_id, new_id, g.user)
    return jsonify(result), status


@app.route("/tickets/<int:ticket_id>/update", methods=["POST"])
@admin_required
def update_ticket(ticket_id):
    db = get_db()
    ticket = db.execute("SELECT * FROM tickets WHERE id = ?", (ticket_id,)).fetchone()
    if ticket is None:
        abort(404)
    status = request.form.get("status", "Open")
    if status not in {"Open", "In Progress", "Closed"}:
        abort(400)
    response = request.form.get("staff_response", "").strip()[:3000]
    db.execute(
        "UPDATE tickets SET status = ?, staff_response = ?, updated_at = ? WHERE id = ?",
        (status, response, utc_now(), ticket_id),
    )
    db.commit()
    flash(f"Ticket #{ticket_id} updated.", "success")
    return redirect(url_for("index"))


@app.cli.command("create-admin")
@click.option("--email", prompt=True)
@click.option("--name", prompt=True)
@click.password_option()
def create_admin(email, name, password):
    """Create an administrator account without exposing credentials in source."""
    if not valid_email(email):
        raise click.ClickException("Enter a valid administrator email address.")
    if len(password) < 12:
        raise click.ClickException("Use an administrator password with at least 12 characters.")
    try:
        get_db().execute(
            "INSERT INTO users (email, name, password_hash, role, created_at) VALUES (?, ?, ?, 'admin', ?)",
            (email.strip().lower(), name.strip(), generate_password_hash(password), utc_now()),
        )
        get_db().commit()
    except sqlite3.IntegrityError as exc:
        raise click.ClickException("An account already exists for that email.") from exc
    click.echo(f"Administrator account created for {email}.")


with app.app_context():
    init_db()


if __name__ == "__main__":
    app.run(debug=os.getenv("FLASK_DEBUG", "0") == "1")
